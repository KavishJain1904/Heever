"""Build the golden-set labelling spreadsheet from the pre-labelled candidates.

    python scripts/build_label_sheet.py

Writes data/golden_200_labels.xlsx. intent / action are pre-filled from the model's
pre-label and are yours to overwrite; confidence starts blank on purpose, so every
row needs one deliberate entry before it counts as labelled. Dropdowns restrict each
column to valid values. The `intents` sheet has the definitions.

Import with: python scripts/run_experiment.py golden data/golden_200_labels.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
from src import load_taxonomy  # noqa: E402

COLUMNS = [
    ("id", 7), ("stratum", 11), ("customer_text", 60), ("brand_first_reply", 45),
    ("intent", 26), ("action", 12), ("confidence", 11), ("ambiguous", 10), ("note", 30),
    ("prelabel_intent", 24), ("prelabel_action", 12), ("prelabel_rationale", 45),
    ("tweet_id", 11), ("conversation_id", 13),
]


def main() -> int:
    rows = json.loads((REPO_ROOT / "data" / "golden_200_prelabels.json").read_text(encoding="utf-8"))
    taxonomy = load_taxonomy()
    names = [i["name"] for i in taxonomy["intents"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append([c for c, _ in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    edit_fill = PatternFill("solid", fgColor="FFF4CC")
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width

    for r in rows:
        brand = next((t["text"] for t in r["thread"] if t["is_brand"]), "")
        ws.append([r["id"], r["stratum"], r["customer_text"], brand,
                   r["prelabel_intent"], r["prelabel_action"], "", "false", "",
                   r["prelabel_intent"], r["prelabel_action"], r["prelabel_rationale"],
                   str(r["tweet_id"]), str(r["conversation_id"])])
    n = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (4, 5, 6, 7, 8):   # editable columns: intent..note
            row[idx].fill = edit_fill

    # Dropdown lists live on a hidden sheet: an inline list of 13 intents exceeds
    # Excel's 255-character limit for literal validation formulas.
    lists = wb.create_sheet("lists")
    for i, name in enumerate(names, 1):
        lists.cell(i, 1, name)
    lists.sheet_state = "hidden"
    for formula, col in ((f"lists!$A$1:$A${len(names)}", "E"),
                         ('"auto_handle,escalate"', "F"),
                         ('"high,med,low"', "G"),
                         ('"true,false"', "H")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:N{n}"

    ref = wb.create_sheet("intents")
    ref.append(["intent", "default_action", "definition", "not_this"])
    for c in ref[1]:
        c.font = Font(bold=True)
    for i in taxonomy["intents"]:
        ref.append([i["name"], i["default_action"], i["definition"].strip(), " | ".join(i.get("not_this") or [])])
    ref.append([])
    ref.append(["precedence", "", taxonomy["precedence_rule"].strip(), ""])
    for col, w in zip("ABCD", (26, 14, 80, 70)):
        ref.column_dimensions[col].width = w
    for row in ref.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = REPO_ROOT / "data" / "golden_200_labels.xlsx"
    wb.save(out)
    print(f"wrote {out} ({len(rows)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
