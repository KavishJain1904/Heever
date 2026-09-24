"""Run the experiment end to end, in stages. Each stage writes committed artifacts.

    python scripts/run_experiment.py predict     # pipeline over all 200 candidates
    python scripts/run_experiment.py golden <labels.csv>   # labelling-tool export -> data/golden_200.csv
    python scripts/run_experiment.py baselines   # needs data/golden_200.csv
    python scripts/run_experiment.py judge       # reply quality: system vs historical vs "DM us"
    python scripts/run_experiment.py tune        # choose tau_intent on non-random strata
    python scripts/run_experiment.py judge-sheet # 30 drafts for blind human scoring
    python scripts/run_experiment.py judge-import data/judge_human_30.xlsx
    python scripts/run_experiment.py judge-import judge_human_30.csv  # from
        # scripts/build_judge_tool.py's judge_human_30.html, same 30 drafts

`predict` does not need labels, so it runs while labelling is still in progress.

Leakage rule: every conversation that appears in the golden candidates or the pilot
30 is removed from the precedent pool BEFORE the index is built. A message must never
retrieve its own brand reply as "how the brand historically resolved this".
"""
from __future__ import annotations

import csv
import json
import sys
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from src import load_config, load_taxonomy  # noqa: E402

DATA = REPO_ROOT / "data"
RESULTS = REPO_ROOT / "artifacts" / "results"
CANDIDATES = DATA / "golden_200_prelabels.json"
PILOT = DATA / "golden_pilot_30.json"
GOLDEN = DATA / "golden_200.csv"

# Pipeline actions -> the annotator's binary decision. DM_HANDOFF counts as
# escalate: the public tweet is automated, but a human does the actual work in DM,
# which is exactly what the annotator was deciding. REQUEST_INFO stays automated.
ACTION_TO_GOLD = {"auto_send": "auto_handle", "request_info": "auto_handle",
                  "dm_handoff": "escalate", "escalate": "escalate"}


def _candidates() -> list[dict]:
    return json.loads(CANDIDATES.read_text(encoding="utf-8"))


def _patch_embedder(model_name: str) -> None:
    """Load MiniLM once. retrieve() calls featurize.embed per query, which would
    otherwise reload the model 200 times."""
    import numpy as np
    from sentence_transformers import SentenceTransformer

    from src import featurize

    model = SentenceTransformer(model_name)
    lock = threading.Lock()

    def embed(texts, _name=None, batch_size=64):
        with lock:
            vecs = model.encode(list(texts), batch_size=batch_size, convert_to_numpy=True,
                                normalize_embeddings=True)
        return vecs.astype(np.float32)

    featurize.embed = embed


def _classify_many(texts: list[str], taxonomy, config, workers: int = 8) -> list[dict]:
    from src.baselines.llm import classify

    with ThreadPoolExecutor(max_workers=workers) as pool:
        return list(pool.map(lambda t: classify([t], taxonomy, config)[0], texts))


def build_index(config):
    import pandas as pd

    from src import featurize
    from src.distill import build_style_card
    from src.retrieve import RetrievalIndex, extract_resolution_precedents

    threads = pd.read_parquet(REPO_ROOT / config.corpus["threads_parquet"])
    held_out = {c["conversation_id"] for c in _candidates()}
    held_out |= {c["conversation_id"] for c in json.loads(PILOT.read_text(encoding="utf-8"))}
    pool = threads[~threads.conversation_id.isin(held_out)]
    records = pool.to_dict("records")
    for r in records:
        r["turns"] = list(r["turns"])

    precedents = extract_resolution_precedents(records)
    kept, dedupe = featurize.dedupe_near_identical([p["brand_reply_normalised"] for p in precedents])
    precedents = [precedents[i] for i in kept]

    taxonomy = load_taxonomy()
    labels = _classify_many([p["customer_text"] for p in precedents], taxonomy, config)
    for p, lab in zip(precedents, labels):
        p["intent"] = lab["intent"]

    embeddings = featurize.embed([p["customer_text"] for p in precedents], config.retrieval["encoder"])
    bm25 = featurize.build_bm25([p["customer_text"] for p in precedents],
                                config.retrieval["bm25_k1"], config.retrieval["bm25_b"])
    brand_replies = [t["text"] for r in records for t in r["turns"] if t["is_brand"]]
    stats = {
        "threads_total": int(len(threads)),
        "threads_held_out": int(len(threads) - len(pool)),
        "brand_replies_in_pool": len(brand_replies),
        "resolution_precedents": len(precedents),
        "precedent_ratio": len(precedents) / len(brand_replies) if brand_replies else 0.0,
        "dedupe": dedupe,
        "precedent_intents": Counter(p["intent"] for p in precedents),
    }
    return RetrievalIndex(precedents, embeddings, bm25), build_style_card(brand_replies), stats


def _draft_from_raw(raw: str):
    """The generator's reply_draft, whether or not the ladder let it be sent."""
    try:
        return json.loads(raw).get("reply_draft") if raw else None
    except (json.JSONDecodeError, AttributeError):
        return None


def cmd_predict() -> int:
    from src.pipeline import run

    config = load_config()
    _patch_embedder(config.retrieval["encoder"])
    index, style_card, stats = build_index(config)
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / "index_stats.json").write_text(json.dumps({**stats, "style_card": style_card}, indent=2))
    print(json.dumps(stats, indent=2))

    taxonomy = load_taxonomy()

    def classifier(text: str):
        out = _classify_many([text], taxonomy, config, workers=1)[0]
        return out["intent"], out["confidence"]

    cands = _candidates()
    messages = [{"message_id": c["id"], "conversation_id": c["conversation_id"],
                 "text": c["customer_text"], "stratum": c["stratum"]} for c in cands]
    records = run(messages, config, classifier=classifier, index=index, style_card=style_card)

    with open(RESULTS / "predictions.jsonl", "w", encoding="utf-8") as fh:
        for r in sorted(records, key=lambda r: r.message_id):
            fh.write(json.dumps({
                "message_id": r.message_id,
                "stratum": r.stratum.value if r.stratum else None,
                "intent": r.predicted_intent,
                "intent_confidence": r.intent_confidence,
                "action": r.decision.action.value,
                "gold_space_action": ACTION_TO_GOLD[r.decision.action.value],
                "escalation_code": r.decision.escalation_code.value if r.decision.escalation_code else None,
                "policy_layer": r.policy_layer_fired.value,
                "grounding_passed": r.grounding_passed,
                "reply_draft": r.decision.reply_draft,
                "candidate_draft": _draft_from_raw(r.raw_model_output),
                "evidence_ids": [e.evidence_id for e in r.evidence],
            }) + "\n")
    print(f"wrote {len(records)} predictions")
    return 0


def _read_labels(path: str) -> list[dict]:
    """The labelling spreadsheet (.xlsx) or a CSV export. Cells become strings.

    Two xlsx producers exist with different sheet names: build_label_sheet.py's
    "labels" sheet (golden-set) and cmd_judge_sheet's "judge_human" sheet (judge
    validation). Both are the workbook's first sheet, so fall back to the active
    sheet when neither known name is present rather than hardcoding one."""
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        for name in ("labels", "judge_human"):
            if name in wb.sheetnames:
                ws = wb[name]
                break
        else:
            ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        return [{k: ("" if v is None else str(v).strip()) for k, v in zip(header, row)} for row in it]
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def cmd_golden_from_prelabels() -> int:
    """Build golden_200.csv from the gpt-5.4 pre-labels with NO human review pass.

    This is the degraded path, taken under deadline. The reference labels are a
    model's, so `reviewed` is "false", `confidence` is "model" rather than a human
    grade, and `label_source` names the model. Nothing downstream computes on those
    three columns; they exist so the artifact states its own provenance and a reader
    cannot mistake this CSV for a hand-labelled one.

    Consequence, stated here because it is easy to forget once the numbers appear:
    every "accuracy" measured against this file is agreement between gpt-5-nano and
    gpt-5.4, not accuracy against ground truth.
    """
    import datetime

    pre = json.loads((DATA / "golden_200_prelabels.json").read_text(encoding="utf-8"))
    rows = pre if isinstance(pre, list) else list(pre.values())
    today = datetime.date.today().isoformat()
    out = [{
        "message_id": r["id"], "tweet_id": r["tweet_id"], "conversation_id": r["conversation_id"],
        "stratum": r["stratum"],
        "gold_intent": r["prelabel_intent"], "gold_action": r["prelabel_action"],
        "quality_of_historical_reply": "",
        "confidence": "model", "ambiguous": "unknown", "note": "",
        "guideline_version": "v2-prelabelled-NO-HUMAN-REVIEW",
        "label_date": today,
        "prelabel_intent": r["prelabel_intent"], "prelabel_action": r["prelabel_action"],
        "reviewed": "false", "label_source": "gpt-5.4-prelabel",
    } for r in rows]
    with open(GOLDEN, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(out[0]))
        writer.writeheader()
        writer.writerows(out)
    report = {
        "n_labelled": len(out), "human_reviewed_rows": 0,
        "label_source": "gpt-5.4-prelabel",
        "strata": Counter(r["stratum"] for r in out),
        "prelabel_override": {
            "intent_override_rate": 0.0, "action_override_rate": 0.0,
            "note": "0 by construction: gold IS the pre-label. Not a measurement.",
        },
        "gold_intents_random": Counter(r["gold_intent"] for r in out if r["stratum"] == "random"),
        "warning": (
            "No human reviewed any row. Metrics against this file measure agreement "
            "between the classifier and gpt-5.4, not accuracy."
        ),
    }
    (RESULTS / "labelling_report.json").write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    print(json.dumps(report, indent=2, default=str))
    return 0


def cmd_golden(labels_path: str) -> int:
    import datetime

    from src import intent_names

    rows = _read_labels(labels_path)
    valid = set(intent_names())
    for r in rows:
        r.setdefault("auto_handle_vs_escalate", r.get("action", ""))
        if not r["auto_handle_vs_escalate"]:
            r["auto_handle_vs_escalate"] = r.get("action", "")
    # A row counts only once confidence is filled: that is the one field the
    # pre-label never supplies, so it proves a human touched the row.
    done = [r for r in rows if r.get("intent") and r.get("auto_handle_vs_escalate") and r.get("confidence")]
    bad = [r["id"] for r in done if r["intent"] not in valid]
    if bad:
        raise SystemExit(f"rows with an intent not in the frozen taxonomy: {bad}")
    today = datetime.date.today().isoformat()
    out = []
    for r in done:
        out.append({
            "message_id": r["id"], "tweet_id": r["tweet_id"], "conversation_id": r["conversation_id"],
            "stratum": r["stratum"], "gold_intent": r["intent"],
            "gold_action": r["auto_handle_vs_escalate"],
            "quality_of_historical_reply": r.get("quality_of_historical_reply", ""),
            "confidence": r["confidence"], "ambiguous": r.get("ambiguous", "false"), "note": r.get("note", ""),
            "guideline_version": r.get("guideline_version") or "v2-prelabelled",
            "label_date": r.get("label_date") or today,
            "prelabel_intent": r.get("prelabel_intent", ""), "prelabel_action": r.get("prelabel_action", ""),
            "reviewed": "true",
        })
    with open(GOLDEN, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(out[0]))
        writer.writeheader()
        writer.writerows(out)

    intent_changed = sum(1 for r in out if r["prelabel_intent"] and r["gold_intent"] != r["prelabel_intent"])
    action_changed = sum(1 for r in out if r["prelabel_action"] and r["gold_action"] != r["prelabel_action"])
    report = {
        "n_labelled": len(out), "n_exported_rows": len(rows),
        "strata": Counter(r["stratum"] for r in out),
        "prelabel_override": {
            "intent_changed": intent_changed, "intent_override_rate": intent_changed / len(out),
            "action_changed": action_changed, "action_override_rate": action_changed / len(out),
        },
        "gold_intents_random": Counter(r["gold_intent"] for r in out if r["stratum"] == "random"),
        "other_unclear_rate_random": (
            sum(1 for r in out if r["stratum"] == "random" and r["gold_intent"] == "other_unclear")
            / max(1, sum(1 for r in out if r["stratum"] == "random"))
        ),
    }
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / "labelling_report.json").write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))
    return 0


def _intent_scores(y_true, y_pred, config) -> dict:
    from src.evaluate import bootstrap_ci, macro_f1, proportion

    return {
        "accuracy": proportion(sum(t == p for t, p in zip(y_true, y_pred)), len(y_true), "accuracy"),
        "macro_f1": macro_f1(y_true, y_pred),
        "macro_f1_ci": bootstrap_ci(macro_f1, y_true, y_pred,
                                    b=int(config.evaluation["bootstrap_replicates"]),
                                    seed=int(config.evaluation["bootstrap_seed"])),
    }


def cmd_baselines() -> int:
    import numpy as np

    from src.baselines import rules, tfidf, trivial
    from src.evaluate import mcnemar

    config = load_config()
    text_by_id = {c["id"]: c["customer_text"] for c in _candidates()}
    with open(GOLDEN, encoding="utf-8", newline="") as fh:
        gold = list(csv.DictReader(fh))
    preds = {p["message_id"]: p for p in map(json.loads, open(RESULTS / "predictions.jsonl", encoding="utf-8"))}

    # Train on the 100 non-headline rows, test on the 100 random rows. The random
    # stratum is never seen by any fitted baseline, so the comparison is clean.
    train = [g for g in gold if g["stratum"] != "random"]
    test = [g for g in gold if g["stratum"] == "random" and g["message_id"] in preds]
    X_tr, y_tr = [text_by_id[g["message_id"]] for g in train], [g["gold_intent"] for g in train]
    X_te, y_te = [text_by_id[g["message_id"]] for g in test], [g["gold_intent"] for g in test]

    systems: dict[str, list[str]] = {
        "majority": trivial.majority_class(y_tr, X_te),
        "random_prior": trivial.random_by_prior(y_tr, X_te, seed=42),
    }
    fitted = rules.fit(X_tr, y_tr)
    systems["rules"] = rules.predict(fitted, X_te)

    min_class = min(Counter(y_tr).values())
    tuned = {"note": "grid search skipped: a class has <2 training rows"}
    if min_class >= 2:
        tuned = tfidf.tune(X_tr, y_tr, cv=min(5, min_class))
        systems["tfidf_svm"] = list(tuned.pop("estimator").predict(X_te))
    else:
        est = tfidf.build_pipeline(min_df=1).fit(X_tr, y_tr)
        systems["tfidf_svm"] = list(est.predict(X_te))

    # Nearest labelled neighbour in MiniLM space: the "retrieval_nn" rung.
    _patch_embedder(config.retrieval["encoder"])
    from src import featurize
    e_tr = featurize.embed(X_tr, None)
    e_te = featurize.embed(X_te, None)
    systems["retrieval_nn"] = [y_tr[int(np.argmax(e_tr @ v))] for v in e_te]

    systems["heever_llm"] = [preds[g["message_id"]]["intent"] for g in test]

    out = {"n_train": len(train), "n_test": len(test), "rules_coverage": rules.coverage(fitted, X_te),
           "tfidf_tuning": tuned, "intent": {}, "action": {}}
    llm_correct = [p == t for p, t in zip(systems["heever_llm"], y_te)]
    for name, y_hat in systems.items():
        out["intent"][name] = _intent_scores(y_te, y_hat, config)
        if name != "heever_llm":
            correct = [p == t for p, t in zip(y_hat, y_te)]
            b = sum(1 for l, o in zip(llm_correct, correct) if l and not o)
            c = sum(1 for l, o in zip(llm_correct, correct) if o and not l)
            stat, p = mcnemar(b, c)
            out["intent"][name]["mcnemar_vs_heever"] = {"b_heever_only": b, "c_baseline_only": c,
                                                        "statistic": stat, "p_value": p}

    # Auto-handle vs escalate on the same 100 rows.
    from src.evaluate import cohens_kappa, proportion
    a_true = [g["gold_action"] for g in test]
    a_pred = [preds[g["message_id"]]["gold_space_action"] for g in test]
    tp = sum(1 for t, p in zip(a_true, a_pred) if t == "escalate" and p == "escalate")
    fn = sum(1 for t, p in zip(a_true, a_pred) if t == "escalate" and p == "auto_handle")
    fp_auto = sum(1 for t, p in zip(a_true, a_pred) if t == "auto_handle" and p == "escalate")
    bad_auto = fn
    out["action"] = {
        "heever_accuracy": proportion(sum(t == p for t, p in zip(a_true, a_pred)), len(a_true), "action_accuracy"),
        "kappa_vs_gold": cohens_kappa(a_true, a_pred),
        "confusion": {"escalate->escalate": tp, "escalate->auto (bad auto-send)": bad_auto,
                      "auto->escalate (needless escalation)": fp_auto,
                      "auto->auto": len(a_true) - tp - bad_auto - fp_auto},
        "escalate_everything_accuracy": proportion(sum(t == "escalate" for t in a_true), len(a_true), "always_escalate"),
        "automation_rate": sum(p == "auto_handle" for p in a_pred) / len(a_pred),
    }
    (RESULTS / "baselines.json").write_text(json.dumps(out, indent=2, default=str))
    print(json.dumps({k: {n: round(v["macro_f1"], 3) for n, v in out["intent"].items()} if k == "intent" else v
                      for k, v in out.items() if k in ("intent", "action")}, indent=2, default=str))
    return 0


def cmd_judge() -> int:
    from src.baselines.trivial import DM_US_REPLY
    from src.judge import judge_one

    config = load_config()
    cands = {c["id"]: c for c in _candidates()}
    preds = [json.loads(l) for l in open(RESULTS / "predictions.jsonl", encoding="utf-8")]
    # The judge must see what the generator saw. The first judge run passed empty
    # evidence: grounding then failed 92/93 drafts ("no retrieved evidence provided")
    # and brand voice flipped with criterion order on 36/93 ("no style guide"). All
    # three systems get the SAME context per message, so the comparison stays fair.
    decisions = {d["message_id"]: d for d in map(json.loads, open(RESULTS / "decisions.jsonl", encoding="utf-8"))}
    style_card = json.loads((RESULTS / "index_stats.json").read_text()).get("style_card", {})
    style_text = ", ".join(f"{k}={v:.2f}" if isinstance(v, float) else f"{k}={v}"
                           for k, v in style_card.items() if k != "note")

    def context_for(message_id: str, action: str) -> str:
        exemplars = decisions.get(message_id, {}).get("evidence", [])
        lines = [f"[{e['evidence_id']}] customer: {e['customer_text']}\n[{e['evidence_id']}] SpotifyCares: {e['brand_reply']}"
                 for e in exemplars] or ["(no precedent retrieved)"]
        return ("Past SpotifyCares resolutions retrieved for this message:\n" + "\n\n".join(lines)
                + f"\n\nSpotifyCares reply style, measured over its public replies: {style_text}"
                + f"\n\nPROPOSED HANDLING: {action} (auto_send = this reply is posted publicly with no human review)")

    jobs = []
    for p in preds:
        if p["stratum"] != "random":
            continue
        c = cands[p["message_id"]]
        brand = next((t["text"] for t in c["thread"] if t["is_brand"]), None)
        # Judge the candidate draft for every random row, not only the ones the
        # ladder auto-sent: otherwise quality is measured on gate survivors only.
        draft = p.get("reply_draft") or p.get("candidate_draft")
        if draft and p["intent"] != "non_support":
            jobs.append(("heever", p, draft))
            jobs.append(("dm_us_degenerate", p, DM_US_REPLY))
            if brand:
                jobs.append(("historical", p, brand))

    def one(job):
        system, p, draft = job
        v = judge_one(cands[p["message_id"]]["customer_text"], draft,
                      context_for(p["message_id"], "auto_send"), config)
        v.update({"system": system, "message_id": p["message_id"], "draft": draft})
        return v

    with ThreadPoolExecutor(max_workers=6) as pool:
        verdicts = list(pool.map(one, jobs))
    path = REPO_ROOT / config.evaluation["judge"]["verdicts_path_full"]
    with open(path, "w", encoding="utf-8") as fh:
        for v in verdicts:
            fh.write(json.dumps(v) + "\n")
    from scripts.redact_artifacts import redact
    redact(path, REPO_ROOT / config.evaluation["judge"]["verdicts_path"])
    summary = {}
    for s in ("heever", "historical", "dm_us_degenerate"):
        rows = [v for v in verdicts if v["system"] == s]
        summary[s] = {"n": len(rows), "all_six_pass": sum(v["all_six_pass"] for v in rows),
                      "position_disagreements": sum(bool(v["position_bias_disagreements"]) for v in rows)}
    print(json.dumps(summary, indent=2))
    return 0


def cmd_errors() -> int:
    """Every wrong call on the random stratum, grouped by (gold -> predicted), with
    the verbatim tweet. The five failure modes in the report are read off this file,
    not guessed: the brief asks for real examples."""
    with open(GOLDEN, encoding="utf-8", newline="") as fh:
        gold = {g["message_id"]: g for g in csv.DictReader(fh)}
    cands = {c["id"]: c for c in _candidates()}
    preds = [json.loads(l) for l in open(RESULTS / "predictions.jsonl", encoding="utf-8")]
    groups: dict[str, list[dict]] = {}
    for p in preds:
        g = gold.get(p["message_id"])
        if not g or g["stratum"] != "random":
            continue
        for kind, truth, guess in (("intent", g["gold_intent"], p["intent"]),
                                   ("action", g["gold_action"], p["gold_space_action"])):
            if truth != guess:
                groups.setdefault(f"{kind}: {truth} -> {guess}", []).append({
                    "id": p["message_id"], "text": cands[p["message_id"]]["customer_text"],
                    "confidence": p["intent_confidence"], "layer": p["policy_layer"],
                    "draft": p.get("candidate_draft"), "note": g.get("note", ""),
                })
    ordered = dict(sorted(groups.items(), key=lambda kv: -len(kv[1])))
    (RESULTS / "errors_random.json").write_text(json.dumps(ordered, indent=2, ensure_ascii=False), encoding="utf-8")
    for key, rows in ordered.items():
        print(f"\n## {key}  ({len(rows)})")
        for r in rows[:3]:
            print(f"  [{r['id']}] conf={r['confidence']:.2f} {r['layer']}: {r['text'][:160]}")
    return 0


def cmd_tune() -> int:
    """Choose tau_intent on the NON-random strata only, then report it on random.

    Rule: among thresholds, keep those whose bad-auto-send share of automated rows
    is at most 1/(1+R), R = cost_model.bad_autosend_vs_needless_escalation (Elkan).
    Of those, pick the one that automates the most. The random stratum is never
    used to choose, so the headline stays uncontaminated by this choice.
    """
    config = load_config()
    ratio = float(config.cost_model["bad_autosend_vs_needless_escalation"])
    max_bad = 1.0 / (1.0 + ratio)
    with open(GOLDEN, encoding="utf-8", newline="") as fh:
        gold = {g["message_id"]: g for g in csv.DictReader(fh)}
    preds = [json.loads(l) for l in open(RESULTS / "predictions.jsonl", encoding="utf-8")]
    auto_by_default = {i["name"] for i in load_taxonomy()["intents"] if i["default_action"] == "auto_send"}

    def simulate(rows, tau):
        # A row automates only if its intent's playbook allows it AND it clears tau
        # AND no earlier veto fired (safety/compliance/no-precedent/context) AND the
        # grounding validator passed. That last clause matters: rung 6
        # (post_generation_validation) sits BELOW the confidence gate, so a row that
        # died at rung 4 in the real run never reached grounding. Re-simulating it at a
        # lower tau without re-applying rung 6 counts drafts the pipeline would veto.
        # Measured cost of omitting it: 4 phantom auto-sends at tau=0.62, 7 at tau=0.55.
        early = {"safety_veto", "compliance_veto", "no_precedent_gate", "context_gate"}
        auto = [p for p in rows if p["intent"] in auto_by_default and p["intent_confidence"] >= tau
                and p["policy_layer"] not in early and p.get("grounding_passed") is True]
        bad = sum(1 for p in auto if gold[p["message_id"]]["gold_action"] == "escalate")
        return {"tau": tau, "n": len(rows), "automated": len(auto),
                "automation_rate": len(auto) / len(rows) if rows else 0.0,
                "bad_auto_sends": bad, "bad_share_of_automated": bad / len(auto) if auto else 0.0}

    tune_rows = [p for p in preds if p["message_id"] in gold and gold[p["message_id"]]["stratum"] != "random"]
    test_rows = [p for p in preds if p["message_id"] in gold and gold[p["message_id"]]["stratum"] == "random"]
    grid = [round(0.30 + 0.05 * i, 2) for i in range(14)]
    sweep = [simulate(tune_rows, t) for t in grid]
    ok = [s for s in sweep if s["automated"] and s["bad_share_of_automated"] <= max_bad]
    chosen = max(ok, key=lambda s: (s["automated"], s["tau"]))["tau"] if ok else max(grid)
    out = {"cost_ratio": ratio, "max_bad_share": max_bad, "chosen_tau_intent": chosen,
           "tuned_on": "stratified+adversarial+policy_trap", "sweep_tune": sweep,
           "sweep_random_report_only": [simulate(test_rows, t) for t in grid],
           "at_chosen_on_random": simulate(test_rows, chosen)}
    (RESULTS / "tau_tuning.json").write_text(json.dumps(out, indent=2))
    print(json.dumps({k: out[k] for k in ("chosen_tau_intent", "max_bad_share", "at_chosen_on_random")}, indent=2))
    return 0


def cmd_judge_sheet() -> int:
    """30 random-stratum Heever drafts for a human to score on the same six criteria."""
    import random

    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    from src.judge import CRITERIA, CRITERION_TEXT

    config = load_config()
    path = REPO_ROOT / config.evaluation["judge"]["verdicts_path_full"]
    verdicts = [json.loads(l) for l in open(path, encoding="utf-8")]
    pool = [v for v in verdicts if v["system"] == "heever"]
    sample = sorted(random.Random(42).sample(pool, min(30, len(pool))), key=lambda v: v["message_id"])
    cands = {c["id"]: c for c in _candidates()}

    wb = Workbook()
    ws = wb.active
    ws.title = "judge_human"
    ws.append(["message_id", "customer_text", "draft", *CRITERIA])
    for v in sample:
        # The judge's answers are deliberately NOT shown: scoring blind is the point.
        ws.append([v["message_id"], cands[v["message_id"]]["customer_text"], v["draft"], *[""] * len(CRITERIA)])
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:I{len(sample) + 1}")
    for col, w in zip("ABCDEFGHI", (10, 55, 55, 14, 14, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ref = wb.create_sheet("criteria")
    for name in CRITERIA:
        ref.append([name, CRITERION_TEXT[name]])
    out = DATA / "judge_human_30.xlsx"
    wb.save(out)
    print(f"wrote {out} ({len(sample)} drafts)")
    return 0


def cmd_judge_import(sheet: str) -> int:
    """Attach the human's yes/no answers to the matching verdicts as human_criteria.

    Accepts either the .xlsx from `judge-sheet` (sheet "judge_human", columns
    message_id, customer_text, draft, <6 criteria>) or a CSV exported by
    scripts/build_judge_tool.py's judge_human_30.html (columns message_id, <6
    criteria> -- no customer_text/draft columns). Both route through
    `_read_labels` so the two input paths behave identically.
    """
    from src.judge import CRITERIA

    config = load_config()
    path = REPO_ROOT / config.evaluation["judge"]["verdicts_path_full"]
    rows = _read_labels(sheet)
    if not rows:
        raise SystemExit(f"no rows read from {sheet}")
    header = list(rows[0].keys())
    criteria_cols = [c for c in header if c in CRITERIA]
    human = {}
    for rec in rows:
        answers = {k: str(rec[k]).strip().lower() for k in criteria_cols if rec.get(k)}
        if len(answers) == len(criteria_cols):
            human[rec["message_id"]] = {k: v == "yes" for k, v in answers.items()}
    verdicts = [json.loads(l) for l in open(path, encoding="utf-8")]
    n = 0
    for v in verdicts:
        if v["system"] == "heever" and v["message_id"] in human:
            v["human_criteria"] = human[v["message_id"]]
            n += 1
    with open(path, "w", encoding="utf-8") as fh:
        for v in verdicts:
            fh.write(json.dumps(v) + "\n")
    # human_criteria was just added to the full file -- regenerate the redacted
    # file from it so the committed copy doesn't go stale (docs/14 human-vs-judge
    # agreement numbers are read off the redacted file via Tier 1).
    from scripts.redact_artifacts import redact
    redact(path, REPO_ROOT / config.evaluation["judge"]["verdicts_path"])
    from src.evaluate import _judge_validation_table
    table = _judge_validation_table(verdicts)
    (RESULTS / "judge_validation.json").write_text(json.dumps(table, indent=2, default=str))
    print(f"attached human scores to {n} verdicts")
    print(json.dumps(table, indent=2, default=str))
    return 0


def main() -> int:
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "predict":
        return cmd_predict()
    if cmd == "golden":
        return cmd_golden(sys.argv[2])
    if cmd == "golden-from-prelabels":
        return cmd_golden_from_prelabels()
    if cmd == "baselines":
        return cmd_baselines()
    if cmd == "judge":
        return cmd_judge()
    if cmd == "tune":
        return cmd_tune()
    if cmd == "errors":
        return cmd_errors()
    if cmd == "judge-sheet":
        return cmd_judge_sheet()
    if cmd == "judge-import":
        return cmd_judge_import(sys.argv[2])
    print(__doc__)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
